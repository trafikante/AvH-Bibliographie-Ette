from pathlib import Path
from copy import deepcopy
from collections import Counter, defaultdict
import csv
import hashlib
import json
import re
import sys
from docx import Document
import build_probe as probe

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = probe.SOURCE
HEADINGS = {category['label']: (category['id'], None) for category in probe.CATEGORIES}
for category in probe.CATEGORIES:
    for subgroup in category['subcategories']:
        HEADINGS[subgroup['source_label']] = (category['id'], subgroup['id'])
SPANISH = {34, 42, 44, 107, 115, 150, 152, 158, 160, 162, 164, 170, 172, 174, 178, 180, 188, 190, 202, 204, 214, 224, 228, 244, 260, 264, 268, 280, 282, 292, 294, 296, 310, 312, 320, 332, 336, 348, 358, 362, 364, 378, 476, 495, 499, 509, 523}
ENGLISH = {70, 72, 76, 80, 146, 186, 194, 226, 230, 238, 256, 284, 300, 302, 304, 322, 324, 326, 328, 334, 338, 368, 372, 384, 388, 412, 424, 436, 438, 448, 450, 452, 460, 470, 474, 480, 507, 527, 533}
FRENCH = {154, 286, 288, 290, 314, 342, 360, 382}
UNDETERMINED = {89, 91, 93, 95, 97, 99, 101, 103, 105, 111}
MAGAZINES = {216, 250, 278, 340, 352, 414, 418}
WEBPAGES = {254, 270, 272}
HIN_TITLE = 'HiN – Alexander von Humboldt im Netz. Internationale Zeitschrift für Humboldt-Studien'
ISSUES = []

def note_issue(record, code, message):
    if message not in record['notes']:
        record['notes'].append(message)
    ISSUES.append({'id': record['id'], 'source_paragraph': record['source']['paragraph_numbers'][0], 'code': code, 'message': message})

def name_from_display(value):
    value = value.strip()
    if ',' not in value:
        return {'literal': value}
    family, given = value.split(',', 1)
    given = given.strip()
    if family.strip() == 'Humboldt' and given in ['Alexander von', 'Alejandro de']:
        given, particle = given.rsplit(' ', 1)
        return probe.person(family.strip(), given, particle)
    return probe.person(family.strip(), given)

def names(value):
    return [name_from_display(part) for part in value.split(' / ')]

def prose_names(value):
    components = re.split(r',\s+| und ', value.strip())
    result = []
    known = {
        'Ottmar Ette': probe.ETTE, 'Oliver Lubrich': probe.person('Lubrich', 'Oliver'),
        'Vera M. Kutzinski': probe.person('Kutzinski', 'Vera M.'),
        'Ingo Schwarz': probe.person('Schwarz', 'Ingo'), 'Irene Prüfer Leske': probe.person('Prüfer Leske', 'Irene'),
        'Julius Ludwig Ideler': probe.person('Ideler', 'Julius Ludwig'),
        'Claudia Kalscheuer': probe.person('Kalscheuer', 'Claudia'),
        'Héctor Pérez Brignoli': probe.person('Pérez Brignoli', 'Héctor'),
        'María Rosario Martí Marco': probe.person('Martí Marco', 'María Rosario'),
        'José Augusto Bezerra': probe.person('Bezerra', 'José Augusto'),
        'Ingrid Schwamborn': probe.person('Schwamborn', 'Ingrid'),
        'Silvia Kruse': probe.person('Kruse', 'Silvia'), 'Matthias Strobel': probe.person('Strobel', 'Matthias'),
        'Johanna Malcher': probe.person('Malcher', 'Johanna'), 'Frank Arnold': probe.person('Arnold', 'Frank'),
        'Antonio Angel Delgado': probe.person('Delgado', 'Antonio Angel'),
        'J. Bradford Anderson': probe.person('Anderson', 'J. Bradford'),
        'J. Ryan Poynter': probe.person('Poynter', 'J. Ryan'), 'Kenneth Berri': probe.person('Berri', 'Kenneth'),
        'Carmen Götz': probe.person('Götz', 'Carmen'), 'Ulrich Päßler': probe.person('Päßler', 'Ulrich')}
    known.update({'Eberhard Knobloch': probe.person('Knobloch', 'Eberhard'), 'Christian Suckow': probe.person('Suckow', 'Christian'), 'Anja Becker': probe.person('Becker', 'Anja')})
    for component in components:
        if component in known:
            result.append(deepcopy(known[component]))
        else:
            result.append({'literal': component})
    return result

def italic_spans(paragraph):
    spans = []
    position = 0
    for run in paragraph.runs:
        length = len(run.text)
        if run.italic and run.text.strip():
            start = position
            stop = position + length
            if spans and not paragraph.text[spans[-1][1]:start].strip():
                spans[-1] = (spans[-1][0], stop)
            else:
                spans.append((start, stop))
        position += length
    return spans

def quoted_title(value):
    opening = value.index('„')
    depth = 0
    for position in range(opening, len(value)):
        character = value[position]
        if character == '„':
            depth += 1
        elif character == '“':
            depth -= 1
            if depth == 0:
                return value[opening + 1:position], position + 1
    raise ValueError('Titelklammer nicht geschlossen: ' + value)

def blocks(document):
    units = []
    category = None
    subtype = None
    for number, paragraph in enumerate(document.paragraphs, 1):
        value = paragraph.text.strip()
        if value in HEADINGS:
            category, subtype = HEADINGS[value]
            continue
        if not value or category is None:
            continue
        new_entry = re.match(r'(Ette|Humboldt|Berghaus|Kutzinski|Lubrich|Holl|Drews|Ren|Asel),', value)
        if category == 'internet-praesentationen' and not value.startswith('URL:'):
            new_entry = True
        if new_entry:
            units.append({'paragraphs': [number], 'category': category, 'subtype': subtype})
        else:
            assert units and units[-1]['category'] == category, (number, value)
            units[-1]['paragraphs'].append(number)
    return units

def language_for(number, title, category):
    if number in UNDETERMINED:
        return [], {'status': 'undetermined', 'basis': 'Publikationssprachen eines ganzen Herausgeberbandes/Heftes aus dem Titel allein nicht sicher feststellbar.'}
    explicit = {127: ['de', 'en', 'es']}
    special = {166: ['pt'], 246: ['pt'], 380: ['pt', 'en', 'de'], 408: ['ru'], 109: ['zh'], 420: ['zh'], 444: ['zh'], 531: ['zh'], 478: ['gl'], 529: ['pt', 'de'], 539: ['de', 'en']}
    if number == 130:
        return [], {'status': 'undetermined', 'basis': 'Quelle nennt vier Sprachen, benennt die vier Sprachfassungen jedoch nicht. Keine ungesicherte Liste ergänzt.'}
    if number in explicit:
        return explicit[number], {'status': 'explicit', 'basis': 'Dreisprachige Präsentation ausdrücklich beschrieben; Titel benennt Deutsch, Englisch und Spanisch.'}
    if number in special:
        codes = special[number]
    elif number in SPANISH:
        codes = ['es']
    elif number in ENGLISH:
        codes = ['en']
    elif number in FRENCH:
        codes = ['fr']
    else:
        codes = ['de']
    return codes, {'status': 'inferred', 'basis': 'Titel und bibliographischer Nachweis redaktionell sprachlich zugeordnet; keine neue Volltextprüfung.'}

def publication_year(value, number):
    if number in [127, 130]:
        return 1998 if number == 127 else 2009
    clean = value.split('DOI:')[0].split('URL:')[0]
    clean = re.split(r',\s*S\.\s*[\divxIVX]', clean, maxsplit=1)[0]
    if number == 216:
        return 2004
    years = re.findall(r'\b(?:19|20)\d{2}\b', clean)
    assert years, (number, value)
    return int(years[-1])

def publisher_from_tail(value, year):
    value = re.split(r',\s*S\.\s*[\divxIVX]', value, maxsplit=1)[0]
    value = value.split('DOI:')[0].split('URL:')[0]
    matches = list(re.finditer(r':\s+([^:]+?)\s+' + str(year) + r'\b', value))
    if not matches:
        return {}
    match = matches[-1]
    place = value[:match.start()].rsplit('. ', 1)[-1].strip()
    if len(place) > 120 or '„' in place:
        return {}
    return {'publisher': match.group(1).strip(), 'publisher-place': place}

def page_from_citation(value):
    matches = list(re.finditer(r',\s*S\.\s*', value))
    if not matches:
        return None
    remainder = value[matches[-1].end():]
    remainder = re.split(r'\.\s+(?:DOI:|URL:|Audioaufnahme)|; der ', remainder, maxsplit=1)[0]
    return remainder.rstrip('. ').strip()

def first_container_title(paragraph, after):
    value = paragraph.text
    eligible = [(start, stop) for start, stop in italic_spans(paragraph) if start >= after]
    if eligible:
        start, stop = eligible[0]
        return value[start:stop].strip().rstrip('.')
    return None

def populate_container(record, paragraph, after, remainder):
    metadata = record['bibliographic']
    number = record['source']['paragraph_numbers'][0]
    title = first_container_title(paragraph, after)
    if 'HiN – Alexander von Humboldt im Netz' in remainder:
        title = HIN_TITLE
    if number in [194, 306]:
        title = 'Northeastern Naturalist' if number == 194 else 'Etudes Germaniques'
    if title:
        metadata['container-title'] = title
    else:
        note_issue(record, 'container-title', 'Übergeordneter Titel noch separat zu prüfen; vollständiger Nachweis ist erhalten.')
    prefix = remainder.split(': ', 1)[0]
    if '(Hg.)' in prefix:
        people = prefix.split(' (Hg.)')[0]
        if people in ['dies.', 'ders.']:
            record['container_contributors']['editor'] = deepcopy(metadata.get('author', []))
        elif ',' in people or len(people) < 130:
            record['container_contributors']['editor'] = names(people)
    elif re.match(r'^[\wÀ-ž -]+, [^:]+$', prefix) and ': ' in remainder:
        record['container_contributors']['author'] = names(prefix)
    if 'et al.' in prefix:
        note_issue(record, 'abbreviated-contributors', 'Herausgeberliste ist in der Quelle mit et al. abgekürzt; keine fehlenden Personen ergänzt.')

def parse_roles(value):
    roles = {}
    expressions = [
        ('editor', r'(?:Herausgegeben(?: und mit (?:einer Einleitung|einem Essay) versehen)?|[Ee]diert und mit einem Nachwort versehen|herausgegeben) von (.+?)(?<![A-Z])\.(?:\s|$)'),
        ('translator', r'(?:Übersetzt|Neu ins Deutsche übersetzt|Aus dem Französischen) von (.+?)(?<![A-Z])\.(?:\s|$)'),
        ('narrator', r'Gelesen von (.+?)(?<![A-Z])\.(?:\s|$)')]
    for role, expression in expressions:
        match = re.search(expression, value)
        if match:
            roles[role] = prose_names(match.group(1))
    if 'aus dem Französischen übersetzt und kommentiert von Ottmar Ette' in value:
        roles['editor'] = [deepcopy(probe.ETTE)]
        roles['translator'] = [deepcopy(probe.ETTE)]
    if 'Nach der Übersetzung aus dem Französischen von Julius Ludwig Ideler' in value:
        roles['translator'] = [probe.person('Ideler', 'Julius Ludwig')]
        roles['editor'] = [deepcopy(probe.ETTE)]
    if 'Übersetzt und herausgegeben von María Rosario Martí Marco und Irene Prüfer Leske' in value:
        assigned = [probe.person('Martí Marco', 'María Rosario'), probe.person('Prüfer Leske', 'Irene')]
        roles['translator'] = assigned
        roles['editor'] = deepcopy(assigned)
    if 'Herausgegeben und neu übersetzt von Irene Prüfer Leske' in value:
        roles['editor'] = [probe.person('Prüfer Leske', 'Irene')]
        roles['translator'] = deepcopy(roles['editor'])
    return roles

def parse_record(unit, document, identifier):
    numbers = unit['paragraphs']
    paragraph = document.paragraphs[numbers[0] - 1]
    value = paragraph.text
    number = numbers[0]
    texts = [document.paragraphs[item - 1].text for item in numbers]
    metadata = {'id': identifier}
    title_end = 0
    if unit['category'] == 'internet-praesentationen':
        metadata['title'] = value.split('. ', 1)[0]
        metadata['type'] = 'webpage'
    else:
        prefix, body = value.split(': ', 1)
        is_editor = prefix.endswith(' (Hg.)')
        if is_editor:
            prefix = prefix.removesuffix(' (Hg.)')
        metadata['editor' if is_editor else 'author'] = names(prefix)
        if body.startswith('„'):
            metadata['title'], title_end = quoted_title(value)
        else:
            spans = italic_spans(paragraph)
            if spans:
                start, title_end = spans[0]
                metadata['title'] = value[start:title_end].strip().rstrip('.')
            else:
                metadata['title'] = body.split('. ', 1)[0].rstrip('.')
        metadata['type'] = 'book'
    year = publication_year(value, number)
    metadata['issued'] = {'date-parts': [[year]]}
    languages, assessment = language_for(number, metadata['title'], unit['category'])
    if len(languages) == 1:
        metadata['language'] = languages[0]
    record = {'id': identifier, 'category_id': unit['category'], 'subcategory_id': unit['subtype'],
        'languages': languages, 'language_assessment': assessment, 'publication_year': year,
        'year_basis': 'online-start' if unit['category'] == 'internet-praesentationen' else 'publication',
        'publication_medium': 'website' if unit['category'] == 'internet-praesentationen' else 'audio' if unit['subtype'] == 'hoerbuecher' or number == 200 else 'text',
        'bibliographic': metadata, 'parallel_titles': [], 'container_contributors': {}, 'related_editions': [], 'relations': [], 'notes': [], 'review_status': 'draft',
        'source': {'file': SOURCE.name, 'sha256': probe.EXPECTED_SHA256, 'paragraph_numbers': numbers, 'paragraphs': texts, 'citation_text': '\n'.join(texts), 'citation_html': '<br>'.join(probe.source_html(document.paragraphs[item - 1]) for item in numbers)}}
    tail = value[title_end:]
    remainder = tail.split('In:', 1)[1].strip() if 'In:' in tail else ''
    publisher = publisher_from_tail(remainder if remainder else tail, year)
    if remainder:
        metadata['type'] = 'chapter' if publisher or number in [134, 150, 254, 370] else 'article-journal'
        populate_container(record, paragraph, title_end, remainder)
        record['container_contributors'].update(parse_roles(remainder))
    else:
        metadata.update(parse_roles(tail))
    if unit['category'] == 'internet-praesentationen':
        metadata['type'] = 'webpage'
    if number in WEBPAGES:
        metadata['type'] = 'webpage'
    if number in MAGAZINES:
        metadata['type'] = 'article-magazine'
    if number == 200:
        metadata['type'] = 'broadcast'
        metadata['medium'] = 'Audioaufnahme des Vortrags'
    if number == 548:
        metadata['type'] = 'article-newspaper'
    if number in [93, 95, 107, 109]:
        metadata['type'] = 'periodical'
    if number in [194, 306, 280]:
        metadata['type'] = 'article-journal'
    if unit['subtype'] == 'hoerbuecher':
        metadata['type'] = 'book'
        metadata['medium'] = 'MP3-CD' if number == 122 else 'Audio-CD'
        metadata['dimensions'] = '3 MP3-CDs' if number == 122 else '2 Audio-CDs'
    metadata.update(publisher)
    page = page_from_citation(value)
    if page:
        metadata['page'] = page
    for address in probe.urls(value):
        if address.startswith('https://doi.org/'):
            metadata['DOI'] = address.removeprefix('https://doi.org/')
        else:
            metadata['URL'] = address
    if number == 127:
        metadata['URL'] = probe.urls(texts[1])[0]
    extent = re.search(r'\[([^\]]+)\s+S\.\]', value)
    if extent:
        metadata['number-of-pages'] = extent.group(1).strip()
    if remainder and metadata['type'] in ['article-journal', 'article-magazine', 'article-newspaper', 'periodical']:
        date_prefix = remainder.split(', S.')[0].split(', S.\u00a0')[0]
        locator = re.search(r'\)\s+([IVXLCDM]+|\d+)(?:,\s*([\d–-]+))?\s*\(', date_prefix)
        if locator:
            if locator.group(2):
                metadata['volume'], metadata['issue'] = locator.group(1), locator.group(2)
            else:
                metadata['issue'] = locator.group(1)
        if number in [146, 152, 162, 164, 166, 170, 172, 186, 198, 222, 230, 232, 238, 244, 246, 248, 252, 264, 280, 284, 290, 302, 306, 316, 326, 330, 332, 344, 348, 358, 364, 378, 382, 384, 408, 410, 414, 416, 442, 456, 458, 460, 464, 466, 474, 487, 525, 527]:
            if locator and not locator.group(2):
                metadata['volume'] = metadata.pop('issue')
    if number in [93, 95]:
        metadata['title'] = value.split(': ', 1)[1].split(' Bd. ', 1)[0].rstrip('.')
        metadata['container-title'] = HIN_TITLE
        metadata['volume'] = 'X' if number == 93 else 'XI'
        metadata['issue'] = '19' if number == 93 else '20'
    if number == 107:
        metadata['title'] = 'Alexander von Humboldt y América Latina'
        metadata['container-title'] = 'Iberoamericana. América Latina – España – Portugal'
        metadata['volume'], metadata['issue'] = 'XIX', '70'
    if number == 130:
        metadata['title'] = 'avhumboldt.de – Alexander von Humboldt Informationen online'
        metadata['editor'] = [deepcopy(probe.ETTE)]
        record['notes'].append('Redaktionsleitung und Konzept: Tobias Kraft, ausdrücklich im Quellenabsatz genannt.')
    if number == 555:
        metadata['type'] = 'chapter'
        record['notes'].append('Übersetzungsrolle von Ette aus der redaktionellen Kategorie erschließbar; im Einzelabsatz nicht explizit benannt.')
    if 'mit einem Diskussionsbeitrag' in value:
        note_issue(record, 'additional-pages', 'Zusätzlicher Diskussionsbeitrag auf S. 80; Hauptseiten 241–252 separat erfasst.')
        metadata['page'] = '241–252'
    if number in [525, 527]:
        metadata['page'] = '16' if number == 525 else '24'
        record['notes'].append('Gesamtbeitrag mit Rede: S. 15–22.' if number == 525 else 'Gesamtbeitrag mit Rede: S. 23–30.')
    if number == 487:
        metadata['number-of-pages'] = '2'
        record['notes'].append('Publish.UP-DOI bezeichnet das Gesamt-Heft; kein Artikel-DOI.')
    if number == 180:
        metadata['container-title'] = 'Debate y Perspectivas'
    if number == 523:
        metadata['container-title'] = 'Iberoamericana. América Latina – España – Portugal'
        metadata['volume'], metadata['issue'] = 'XIX', '70'
    if number == 68:
        metadata['title'] = value.split(': ', 1)[1].split('. Frankfurt am Main / Leipzig:', 1)[0]
    if number in [266, 501]:
        metadata['container-title'] = remainder.split('Humboldt, Alexander von: ', 1)[1].split('. Frankfurt am Main / Leipzig:', 1)[0]
    if number in [362, 368, 372, 422, 432, 456]:
        field = 'volume' if number in [372, 422] else 'issue'
        for key in ['volume', 'issue']:
            metadata.pop(key, None)
        metadata[field] = {362: '829–830', 368: '128', 372: '45', 422: '48', 432: '26', 456: '194'}[number]
    if number == 370:
        metadata.update({'publisher-place': 'Berlin', 'publisher': 'Berlin-Brandenburgische Akademie der Wissenschaften', 'version': '1'})
        record['container_contributors']['editor'] = [deepcopy(probe.ETTE)]
    if number == 134:
        metadata['publisher-place'] = 'Berlin'
        record['notes'].append('Verlag ist im Quellenabsatz nicht angegeben.')
    count = re.search(r'(\d+)\s+Bde\.', value)
    if count:
        metadata['number-of-volumes'] = count.group(1)
    if 'Sondernummer' in value or 'Special Issue' in value:
        record['notes'].append('Sonderheft-/Dossiertitel und Kontext sind in der vollständigen Quellenangabe erhalten.')
    if metadata['type'] == 'chapter' and 'publisher' not in metadata:
        note_issue(record, 'missing-imprint', 'Impressum nicht vollständig separat erschlossen; vollständige Quellenangabe bleibt erhalten.')
    return record

def append_editions(record, document):
    primary = record['source']['paragraph_numbers'][0]
    for number in record['source']['paragraph_numbers'][1:]:
        if primary == 127:
            continue
        value = document.paragraphs[number - 1].text
        if number == 55:
            editions = [(2004, 'Zweite Auflage'), (2004, 'Dritte Auflage'), (2004, 'Sonderausgabe (2004/2006)'), (2014, 'Nachdruck der Ausgabe von 2004, Juli 2014')]
        elif number == 58:
            editions = [(2004, 'Zweite Auflage'), (2004, 'Dritte Auflage'), (2004, 'Sonderausgabe (2004/2006)')]
        else:
            years = re.findall(r'\b(?:19|20)\d{2}\b', value.split('URL:')[0])
            assert years, (number, value)
            editions = [(int(years[0]), value.split(':', 1)[0].split('. ', 1)[0])]
        for year, label in editions:
            entry = {'publication_year': year, 'label': label, 'source_paragraph': number, 'citation_text': value}
            if '2004/2006' in label:
                entry['publication_years'] = [2004, 2006]
            if probe.urls(value):
                entry['url'] = probe.urls(value)[0]
            record['related_editions'].append(entry)

def dataset_schema():
    schema = probe.build_schema()
    schema['title'] = 'AvH-Bibliographie-Ette vollständiger Erfassungsstand'
    schema['properties']['schema_version']['const'] = '0.2.0'
    schema['properties']['dataset']['properties']['status']['const'] = 'complete-transcription'
    schema['properties']['dataset']['properties']['selection_principles'] = {'type': 'array', 'items': {'type': 'string'}}
    item = schema['properties']['records']['items']
    item['required'].append('source_order')
    item['properties']['source_order'] = {'type': 'integer', 'minimum': 1}
    item['properties']['languages']['items']['enum'].append('gl')
    item['properties']['parallel_titles']['items']['properties']['language']['enum'].append('gl')
    bibliographic = item['properties']['bibliographic']
    bibliographic['properties']['type']['enum'].extend(['article-magazine', 'broadcast'])
    bibliographic['properties']['narrator'] = deepcopy(bibliographic['properties']['author'])
    bibliographic['properties']['version'] = {'type': 'string'}
    for role in ['author', 'editor', 'translator', 'narrator']:
        bibliographic['properties'][role]['items'] = {'type': 'object', 'properties': {'family': {'type': 'string'}, 'given': {'type': 'string'}, 'non-dropping-particle': {'type': 'string'}, 'literal': {'type': 'string'}}, 'additionalProperties': False}
    for role in item['properties']['container_contributors']['properties'].values():
        role['items'] = deepcopy(bibliographic['properties']['author']['items'])
    item['properties']['related_editions']['items']['properties']['publication_years'] = {'type': 'array', 'items': {'type': 'integer'}}
    return schema

def write_table(records):
    category_labels = {category['id']: category['label'] for category in probe.CATEGORIES}
    subcategory_labels = {subgroup['id']: subgroup['label'] for category in probe.CATEGORIES for subgroup in category['subcategories']}
    columns = ['id', 'quellreihenfolge', 'kategorie', 'untergruppe', 'publikationstyp_csl', 'jahr', 'jahresgrundlage', 'sprachen', 'sprachzuordnung', 'titel', 'personen_mit_rollen', 'uebergeordneter_titel', 'personen_des_uebergeordneten_werks', 'verlag', 'verlagsort', 'band', 'heft', 'seiten', 'doi', 'url', 'weitere_ausgaben_jahre', 'pruefstatus', 'quellabsaetze', 'literaturangabe', 'online_verweise', 'weitere_mitwirkende']
    with (PROJECT / 'data/bibliographie.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for record in records:
            metadata = record['bibliographic']
            roles = {role: metadata[role] for role in ['author', 'editor', 'translator', 'narrator'] if role in metadata}
            years = sorted({year for edition in record['related_editions'] for year in edition.get('publication_years', [edition['publication_year']])})
            writer.writerow(dict(zip(columns, [record['id'], record['source_order'], category_labels[record['category_id']], subcategory_labels.get(record['subcategory_id'], ''), metadata['type'], record['publication_year'], record['year_basis'], ';'.join(record['languages']), record['language_assessment']['status'], metadata['title'], json.dumps(roles, ensure_ascii=False), metadata.get('container-title', ''), json.dumps(record['container_contributors'], ensure_ascii=False), metadata.get('publisher', ''), metadata.get('publisher-place', ''), metadata.get('volume', ''), metadata.get('issue', ''), metadata.get('page', ''), metadata.get('DOI', ''), metadata.get('URL', ''), ';'.join(map(str, years)), record['review_status'], ';'.join(map(str, record['source']['paragraph_numbers'])), record['source']['citation_text'], json.dumps(record.get('online_resources', []), ensure_ascii=False), json.dumps(record.get('contributors', []), ensure_ascii=False)])))
    with (PROJECT / 'data/bibliographie.csv').open(encoding='utf-8-sig', newline='') as handle:
        reloaded = list(csv.DictReader(handle))
    assert len(reloaded) == len(records)
    assert [row['literaturangabe'] for row in reloaded] == [record['source']['citation_text'] for record in records]

def validate(dataset, schema, document, units, sample):
    probe.validate_instance(dataset, schema)
    records = dataset['records']
    assert len(records) == len(units) == dataset['dataset']['record_count']
    identifiers = {record['id'] for record in records}
    assert len(identifiers) == len(records)
    expected_paragraphs = [number for unit in units for number in unit['paragraphs']]
    found_paragraphs = [number for record in records for number in record['source']['paragraph_numbers']]
    assert expected_paragraphs == found_paragraphs
    assert len(found_paragraphs) == len(set(found_paragraphs))
    assert [record['source_order'] for record in records] == list(range(1, len(records) + 1))
    for record in records:
        metadata = record['bibliographic']
        assert metadata['id'] == record['id']
        assert metadata['issued']['date-parts'][0][0] == record['publication_year']
        assert metadata['title'].strip(), record['id']
        expected = [document.paragraphs[number - 1].text for number in record['source']['paragraph_numbers']]
        assert expected == record['source']['paragraphs']
        assert '\n'.join(expected) == record['source']['citation_text']
        for relation in record['relations']:
            assert relation['target_id'] in identifiers
        for role in ['author', 'editor', 'translator', 'narrator']:
            for name in metadata.get(role, []):
                assert bool(name.get('literal')) != bool(name.get('family') and name.get('given')), name
    source_lookup = {record['source']['paragraph_numbers'][0]: record for record in records}
    for sample_record in sample:
        matching = source_lookup[sample_record['source']['paragraph_numbers'][0]]
        assert matching['id'] == sample_record['id']
        assert matching['source'] == sample_record['source']
    assert source_lookup[84]['publication_year'] == 2024
    assert source_lookup[430]['publication_year'] == 2019
    assert source_lookup[529]['publication_year'] == 2019
    assert source_lookup[539]['languages'] == ['de', 'en']
    assert source_lookup[478]['languages'] == ['gl']
    assert source_lookup[380]['languages'] == ['pt', 'en', 'de']
    assert source_lookup[130]['languages'] == []
    assert source_lookup[525]['bibliographic']['page'] == '16'
    assert source_lookup[527]['bibliographic']['page'] == '24'
    assert source_lookup[408]['bibliographic']['title'].count('ʹ') == 2

def main():
    original = SOURCE.read_bytes()
    assert hashlib.sha256(original).hexdigest() == probe.EXPECTED_SHA256, 'Quellenfassung geändert; Zuordnungen erneut prüfen.'
    document = Document(SOURCE)
    units = blocks(document)
    sample = json.loads((PROJECT / 'data/probe.json').read_text(encoding='utf-8'))['records']
    sample_lookup = {record['source']['paragraph_numbers'][0]: record for record in sample}
    next_id = 21
    records = []
    for source_order, unit in enumerate(units, 1):
        number = unit['paragraphs'][0]
        if number in sample_lookup:
            record = deepcopy(sample_lookup[number])
            assert record['source']['paragraph_numbers'] == unit['paragraphs']
        else:
            record = parse_record(unit, document, f'ette-{next_id:06d}')
            append_editions(record, document)
            next_id += 1
        record['source_order'] = source_order
        records.append(record)
    lookup = {record['source']['paragraph_numbers'][0]: record for record in records}
    for record in records:
        number = record['source']['paragraph_numbers'][0]
        if number in [380, 478, 539]:
            record['languages'], record['language_assessment'] = language_for(number, record['bibliographic']['title'], record['category_id'])
        if not record['languages']:
            note_issue(record, 'publication-languages', record['language_assessment']['basis'])
    lookup[78]['relations'].append({'type': 'same-publication', 'target_id': lookup[557]['id']})
    lookup[557]['relations'].append({'type': 'same-publication', 'target_id': lookup[78]['id']})
    for original_number, child_numbers in [(50, [138, 140]), (54, [218, 489]), (57, [491]), (60, [220, 493]), (63, [240, 497, 555]), (66, []), (68, [266, 501]), (70, [304]), (72, [334]), (74, [386]), (76, [388]), (78, [390, 392, 517, 519]), (80, [424, 533]), (82, [462]), (84, [472]), (89, [182, 184]), (91, [192, 485]), (99, [366, 511]), (101, [374, 513]), (103, [376, 515]), (105, [396, 398, 400, 402, 416, 521]), (107, [523]), (111, [468, 543])]:
        for child_number in child_numbers:
            if child_number == 416:
                continue
            relationship = {'type': 'part-of', 'target_id': lookup[original_number]['id']}
            if relationship not in lookup[child_number]['relations']:
                lookup[child_number]['relations'].append(relationship)
    dataset = {'schema_version': '0.2.0', 'dataset': {'title': 'AvH-Bibliographie-Ette', 'status': 'complete-transcription', 'created': probe.DATE, 'record_count': len(records), 'source_file': SOURCE.name, 'source_sha256': probe.EXPECTED_SHA256, 'selection_principles': [document.paragraphs[number - 1].text for number in [7, 11, 12, 13]]}, 'categories': deepcopy(probe.CATEGORIES), 'records': records}
    schema = dataset_schema()
    validate(dataset, schema, document, units, sample)
    if '--check' in sys.argv:
        existing = json.loads((PROJECT / 'data/bibliographie.json').read_text(encoding='utf-8'))
        validate(existing, schema, document, units, sample)
        assert existing == dataset
        assert json.loads((PROJECT / 'schema/bibliographie.schema.json').read_text(encoding='utf-8')) == schema
        print('Vollständigkeits- und Strukturprüfung erfolgreich: ' + str(len(records)) + ' Einträge; Probe-IDs erhalten; Quelle unverändert.')
        return
    (PROJECT / 'data/bibliographie.json').write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (PROJECT / 'schema/bibliographie.schema.json').write_text(json.dumps(schema, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    write_table(records)
    (PROJECT / 'data/erfassungsfragen.json').write_text(json.dumps(ISSUES, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    assert SOURCE.read_bytes() == original
    summary = {'records': len(records), 'source_paragraphs': sum(len(unit['paragraphs']) for unit in units), 'categories': dict(Counter(record['category_id'] for record in records)), 'subcategories': dict(Counter(record['subcategory_id'] for record in records if record['subcategory_id'])), 'types': dict(Counter(record['bibliographic']['type'] for record in records)), 'language_assessments': dict(Counter(record['language_assessment']['status'] for record in records)), 'language_occurrences': dict(Counter(code for record in records for code in record['languages'])), 'issues': ISSUES, 'source_unchanged': True}
    (PROJECT / 'data/erfassungsstatistik.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
