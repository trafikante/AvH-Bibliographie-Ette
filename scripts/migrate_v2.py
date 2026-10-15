from pathlib import Path
from copy import deepcopy
from collections import Counter
from difflib import SequenceMatcher
from datetime import date
import csv
import hashlib
import json
import re
import sys
from docx import Document
import build_dataset as legacy
import build_probe as probe

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT.parent / 'bibliographie_ette_final_v2.docx'
BASELINE = PROJECT / 'data/erfassungsstand_vor_v2.json'


def read_units(document):
    headings = dict(legacy.HEADINGS)
    headings['Wissenschaftskommunikation'] = ('debatten', None)
    units = []
    category = None
    subgroup = None
    for number, paragraph in enumerate(document.paragraphs, 1):
        text = paragraph.text.strip()
        if text in headings:
            category, subgroup = headings[text]
        elif text and category:
            starts = re.match(r'^(Ette|Humboldt|Berghaus|Kutzinski|Lubrich|Holl|Drews|Ren|Asel|Wedel|Merseburger|Dalheimer),', text)
            if starts or category == 'internet-praesentationen' and not text.startswith('URL:'):
                units.append({'paragraphs': [number], 'category': category, 'subtype': subgroup})
            else:
                assert units and units[-1]['category'] == category, (number, text)
                units[-1]['paragraphs'].append(number)
    return units


def source_for(unit, document, digest):
    paragraphs = [document.paragraphs[number - 1] for number in unit['paragraphs']]
    texts = [paragraph.text for paragraph in paragraphs]
    return {'file': SOURCE.name, 'sha256': digest, 'paragraph_numbers': unit['paragraphs'],
            'paragraphs': texts, 'citation_text': '\n'.join(texts),
            'citation_html': '<br>'.join(probe.source_html(paragraph) for paragraph in paragraphs)}


def media_record(unit, document, identifier, source):
    paragraph = document.paragraphs[unit['paragraphs'][0] - 1]
    text = paragraph.text
    prefix, body = text.split(': ', 1)
    prefix = prefix.removesuffix(' (Mitwirkung)')
    if body.startswith('„'):
        title, end = legacy.quoted_title(text)
    else:
        match = re.match(r'\[([^]]+)\]', body)
        assert match, text
        title = '[' + match.group(1) + ']'
        end = text.index(']') + 1
    remainder = text[end:].split('In: ', 1)[1]
    dates = re.findall(r'(?<!\d)(\d{1,2})\.(\d{1,2})\.((?:19|20)\d{2})', remainder.split('URL:')[0])
    if dates:
        day, month, year = map(int, dates[0])
        issued = [year, month, day]
    else:
        year = int(re.findall(r'\b(?:19|20)\d{2}\b', remainder.split('URL:')[0])[0])
        issued = [year]
    video = any(word in text for word in ['Fernseh', 'Video'])
    metadata = {'id': identifier, 'type': 'broadcast', 'title': title,
                'author': legacy.names(prefix), 'issued': {'date-parts': [issued]},
                'language': 'en' if 'Englischsprachiges' in text else 'de',
                'container-title': legacy.first_container_title(paragraph, end) or remainder.split(' (', 1)[0],
                'medium': 'Video' if video else 'Audio'}
    urls = probe.urls(text)
    if urls:
        metadata['URL'] = urls[0]
    contributors = []
    if '(Mitwirkung)' in text or 'Mitwirkung: Ottmar Ette' in text:
        contributors.append({'role': 'participant', 'name': deepcopy(probe.ETTE)})
        if '(Mitwirkung)' in text:
            metadata.pop('author', None)
    for marker, role in [('Gespräch mit ', 'interviewer'), ('Moderation: ', 'moderator')]:
        if marker in text:
            people = text.split(marker, 1)[1].split('. In:', 1)[0]
            for person in legacy.prose_names(people):
                contributors.append({'role': role, 'name': person})
    return {'id': identifier, 'category_id': unit['category'], 'subcategory_id': unit['subtype'],
            'languages': [metadata['language']],
            'language_assessment': {'status': 'explicit' if 'Englischsprachiges' in text else 'inferred',
                                   'basis': 'Sprachangabe bzw. Titel und Sendungskontext der aktuellen Quelle; keine erneute Hörprüfung.'},
            'publication_year': year, 'year_basis': 'publication',
            'publication_medium': 'video' if video else 'audio', 'bibliographic': metadata,
            'parallel_titles': [], 'container_contributors': {}, 'related_editions': [],
            'relations': [], 'contributors': contributors, 'notes': [],
            'review_status': 'draft', 'source': source}


def main():
    check = '--check' in sys.argv
    if not BASELINE.exists():
        assert not check
        BASELINE.write_bytes((PROJECT / 'data/bibliographie.json').read_bytes())
    baseline = json.loads(BASELINE.read_text(encoding='utf-8'))
    decisions = json.loads((PROJECT / 'data/redaktionelle-festlegungen.json').read_text(encoding='utf-8'))
    original = SOURCE.read_bytes()
    digest = hashlib.sha256(original).hexdigest()
    document = Document(SOURCE)
    units = read_units(document)
    old_records = baseline['records']
    next_id = max(int(record['id'].split('-')[1]) for record in old_records) + 1
    used = set()
    records = []
    changes = []
    for order, unit in enumerate(units, 1):
        source = source_for(unit, document, digest)
        text = source['citation_text']
        candidates = [record for record in old_records if record['id'] not in used and record['category_id'] == unit['category']]
        exact = [record for record in candidates if record['source']['citation_text'] == text]
        match = exact[0] if len(exact) == 1 else None
        if match is None and unit['category'] == 'debatten' and 'Wissen gesellschaftsfähig machen' in text:
            moved = [record for record in old_records if record['id'] not in used and record['bibliographic']['title'] == 'Wissen gesellschaftsfähig machen – Alexander von Humboldt und die Informationsgesellschaft']
            assert len(moved) == 1
            match = moved[0]
        if match is None and candidates:
            scored = sorted([(SequenceMatcher(None, text, record['source']['citation_text']).ratio(), record) for record in candidates], key=lambda pair: pair[0], reverse=True)
            if scored[0][0] >= 0.83 and (len(scored) == 1 or scored[0][0] - scored[1][0] > 0.05):
                match = scored[0][1]
        if match:
            used.add(match['id'])
            record = deepcopy(match)
            if not exact:
                changes.append({'id': match['id'], 'kind': 'changed', 'before': match['source']['citation_text'], 'after': text})
                if unit['category'] == 'debatten' and 'Audioaufnahme' in text:
                    record = media_record(unit, document, match['id'], source)
                else:
                    old_numbers = match['source']['paragraph_numbers']
                    proxy = type('SourceView', (), {})()
                    proxy.paragraphs = list(document.paragraphs)
                    while len(proxy.paragraphs) < max(old_numbers):
                        proxy.paragraphs.append(document.paragraphs[0])
                    for old_number, new_number in zip(old_numbers, unit['paragraphs']):
                        proxy.paragraphs[old_number - 1] = document.paragraphs[new_number - 1]
                    record = legacy.parse_record({'paragraphs': old_numbers, 'category': unit['category'], 'subtype': unit['subtype']}, proxy, match['id'])
                    legacy.append_editions(record, proxy)
                    record['relations'] = deepcopy(match['relations'])
            record['source'] = source
            old_numbers = match['source']['paragraph_numbers']
            for edition in record['related_editions']:
                if edition['source_paragraph'] in old_numbers:
                    edition['source_paragraph'] = unit['paragraphs'][old_numbers.index(edition['source_paragraph'])]
        else:
            assert unit['category'] == 'debatten', ('Nicht sicher zugeordneter Eintrag', text)
            identifier = f'ette-{next_id:06d}'
            next_id += 1
            record = media_record(unit, document, identifier, source)
            changes.append({'id': identifier, 'kind': 'added', 'after': text})
        record['source_order'] = order
        record['category_id'] = unit['category']
        record['subcategory_id'] = unit['subtype']
        metadata = record['bibliographic']
        if record['subcategory_id'] == 'sammelbaende' and metadata['type'] == 'periodical':
            metadata['type'] = decisions['dossier_bibliographic_type']
            record['notes'].append('Zeitschriftendossier gemäß Nutzerentscheidung wie Sammelband erfasst; Zeitschrift, Band und Heft bleiben erhalten.')
        if record['id'] in decisions['languages_by_id']:
            record['languages'] = decisions['languages_by_id'][record['id']]
            record['language_assessment'] = {'status': 'explicit', 'basis': 'Publikationssprachen vom Nutzer einzeln bestätigt; redaktionelle Festlegung nach stabiler ID.'}
            metadata.pop('language', None)
            if len(record['languages']) == 1:
                metadata['language'] = record['languages'][0]
        record['online_resources'] = [{'url': address, 'basis': 'current-source'} for address in probe.urls(text)]
        if metadata['type'] in ['webpage', 'article-journal', 'article-magazine', 'article-newspaper', 'broadcast']:
            metadata.pop('publisher', None)
            metadata.pop('publisher-place', None)
        if metadata.get('container-title') == 'Northeastern Naturalist' and metadata.get('issue') == '8':
            metadata['volume'] = metadata.pop('issue')
        records.append(record)
    retired = [record for record in old_records if record['id'] not in used]
    identifiers = {record['id'] for record in records}
    for record in records:
        record['relations'] = [relation for relation in record['relations'] if relation['target_id'] in identifiers]
    categories = deepcopy(probe.CATEGORIES)
    for category in categories:
        if category['id'] == 'debatten':
            category['label'] = 'Wissenschaftskommunikation'
        if category['id'] == 'akademische-buchpublikationen':
            category['label'] = decisions['academic_category_label']
    dataset = {'schema_version': '0.2.0', 'dataset': {'title': 'AvH-Bibliographie-Ette', 'status': 'complete-transcription',
               'created': probe.DATE, 'record_count': len(records), 'source_file': SOURCE.name, 'source_sha256': digest,
               'selection_principles': [document.paragraphs[number - 1].text for number in [7, 11, 12, 13]]},
               'categories': categories, 'records': records}
    dataset['dataset']['license'] = decisions['license']
    dataset['dataset']['license_url'] = decisions['license_url']
    schema = legacy.dataset_schema()
    schema['properties']['dataset']['properties']['license'] = {'const': 'CC-BY-4.0'}
    schema['properties']['dataset']['properties']['license_url'] = {'const': 'https://creativecommons.org/licenses/by/4.0/'}
    schema['properties']['records']['items']['properties']['publication_medium']['enum'].append('video')
    date_schema = schema['properties']['records']['items']['properties']['bibliographic']['properties']['issued']['properties']['date-parts']['items']
    date_schema['maxItems'] = 3
    date_schema['items'] = {'type': 'integer', 'minimum': 1, 'maximum': 2100}
    item_schema = schema['properties']['records']['items']
    item_schema['properties']['online_resources'] = {'type': 'array', 'items': {'type': 'object', 'required': ['url', 'basis'], 'properties': {'url': {'type': 'string'}, 'basis': {'const': 'current-source'}}, 'additionalProperties': False}}
    item_schema['properties']['contributors'] = {'type': 'array', 'items': {'type': 'object', 'required': ['role', 'name'], 'properties': {'role': {'enum': ['participant', 'interviewer', 'moderator']}, 'name': deepcopy(item_schema['properties']['bibliographic']['properties']['author']['items'])}, 'additionalProperties': False}}
    probe.validate_instance(dataset, schema)
    found = [number for record in records for number in record['source']['paragraph_numbers']]
    assert found == [number for unit in units for number in unit['paragraphs']]
    assert len(found) == len(set(found))
    assert len(identifiers) == len(records)
    for record in records:
        issued = record['bibliographic']['issued']['date-parts'][0]
        assert issued[0] == record['publication_year']
        if len(issued) == 3:
            date(*issued)
        assert record['bibliographic']['id'] == record['id']
    assert SOURCE.read_bytes() == original
    issues = [{'id': record['id'], 'code': 'publication-languages', 'message': record['language_assessment']['basis']}
              for record in records if not record['languages']]
    issues += [{'id': record['id'], 'code': 'metadata-review', 'message': note}
               for record in records for note in record['notes']
               if any(marker in note for marker in ['noch separat zu prüfen', 'Impressum nicht vollständig', 'Verlag ist im Quellenabsatz nicht angegeben', 'et al. abgekürzt', 'keine unabhängige Rollenprüfung'])]
    audit = {'source_sha256': digest, 'record_count': len(records), 'unchanged': len(records) - len(changes),
             'changes': changes, 'retired_ids': [record['id'] for record in retired],
             'retired_citations': [record['source']['citation_text'] for record in retired],
             'source_unchanged': True, 'categories': dict(Counter(record['category_id'] for record in records))}
    outputs = {'data/bibliographie.json': dataset, 'schema/bibliographie.schema.json': schema,
               'data/v2-abgleich.json': audit, 'data/erfassungsfragen.json': issues,
               'data/erfassungsstatistik.json': {'records': len(records), 'source_paragraphs': len(found),
               'categories': audit['categories'], 'types': dict(Counter(record['bibliographic']['type'] for record in records)),
               'language_assessments': dict(Counter(record['language_assessment']['status'] for record in records)), 'source_unchanged': True}}
    for filename, value in outputs.items():
        path = PROJECT / filename
        if check:
            assert json.loads(path.read_text(encoding='utf-8')) == value, filename
        else:
            path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    if not check:
        probe.CATEGORIES = categories
        legacy.write_table(records)
    with (PROJECT / 'data/bibliographie.csv').open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.DictReader(handle))
    assert [row['id'] for row in rows] == [record['id'] for record in records]
    assert [row['literaturangabe'] for row in rows] == [record['source']['citation_text'] for record in records]
    assert [json.loads(row['online_verweise']) for row in rows] == [record['online_resources'] for record in records]
    assert all(relation['target_id'] in identifiers for record in records for relation in record['relations'])
    print(json.dumps({key: value for key, value in audit.items() if key not in ['changes', 'retired_citations']}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
