from pathlib import Path
from copy import deepcopy
import csv
import hashlib
import html
import json
import re
import sys
from collections import Counter
from docx import Document

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT.parent / 'bibliographie_ette_final.docx'
EXPECTED_SHA256 = 'd403b5c6d2811644a64e0473778780232961c709d6515118d4b913d982495fd3'
DATE = '2026-10-13'
NAMESPACES = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
CATEGORIES = [
    {'id': 'literarisches', 'label': 'Literarisches', 'subcategories': []},
    {'id': 'akademische-buchpublikationen', 'label': 'Akademische Buchpublikationen', 'subcategories': [
        {'id': 'monographien', 'label': 'Monographien', 'source_label': 'a. Monographien'},
        {'id': 'editionen', 'label': 'Editionen', 'source_label': 'b. Editionen'},
        {'id': 'sammelbaende', 'label': 'Sammelbände', 'source_label': 'c. Sammelbände'},
        {'id': 'didaktische-materialien', 'label': 'Didaktische Materialien', 'source_label': 'd. Didaktische Materialien'},
        {'id': 'hoerbuecher', 'label': 'Hörbücher', 'source_label': 'e. Hörbücher'}]},
    {'id': 'internet-praesentationen', 'label': 'Internet-Präsentationen', 'subcategories': []},
    {'id': 'aufsaetze', 'label': 'Aufsätze', 'subcategories': []},
    {'id': 'vorworte-lexikonartikel-bibliographien', 'label': 'Vorworte, Lexikonartikel und Bibliographien', 'subcategories': []},
    {'id': 'debatten', 'label': 'Debatten', 'subcategories': []},
    {'id': 'uebersetzungen', 'label': 'Übersetzungen', 'subcategories': []}
]
LANGUAGES = {'de': 'Deutsch', 'en': 'Englisch', 'es': 'Spanisch', 'pt': 'Portugiesisch', 'fr': 'Französisch', 'ru': 'Russisch', 'zh': 'Chinesisch'}

def person(family, given, particle=None):
    value = {'family': family, 'given': given}
    if particle:
        value['non-dropping-particle'] = particle
    return value

ETTE = person('Ette', 'Ottmar')
HUMBOLDT = person('Humboldt', 'Alexander', 'von')
ACADEMIC = 'akademische-buchpublikationen'
PREFACES = 'vorworte-lexikonartikel-bibliographien'

def seed(paragraphs, category, subtype, languages, year, standard_type, **fields):
    return {'paragraphs': paragraphs, 'category': category, 'subtype': subtype,
            'languages': languages, 'year': year, 'standard_type': standard_type, 'fields': fields}

SEEDS = [
    seed([27], 'literarisches', None, ['de'], 2025, 'article-journal', author=[ETTE], container_title='HiN – Alexander von Humboldt im Netz. Internationale Zeitschrift für Humboldt-Studien', volume='XXVI', issue='50', page='249–259'),
    seed([34], ACADEMIC, 'monographien', ['es'], 2001, 'book', title='Literatura de viaje: de Humboldt a Baudrillard', author=[ETTE], translator=[person('Delgado', 'Antonio Angel')], publisher='Universidad Nacional Autónoma de México', publisher_place='México, D.F.', collection_title='Colección Jornadas', number_of_pages='122'),
    seed([36, 37, 38], ACADEMIC, 'monographien', ['de'], 2002, 'book', title='Weltbewußtsein. Alexander von Humboldt und das unvollendete Projekt einer anderen Moderne', author=[ETTE], publisher='Velbrück Wissenschaft', publisher_place='Weilerswist', number_of_pages='243'),
    seed([50, 51, 52], ACADEMIC, 'editionen', ['de'], 1991, 'book', title='Reise in die Äquinoktial-Gegenden des Neuen Kontinents', author=[HUMBOLDT], editor=[ETTE], publisher='Insel Verlag', publisher_place='Frankfurt am Main / Leipzig', number_of_volumes='2', number_of_pages='1637'),
    seed([89], ACADEMIC, 'sammelbaende', [], 2001, 'book', title='Ansichten Amerikas. Neuere Studien zu Alexander von Humboldt', editor=[ETTE, person('Bernecker', 'Walther L.')], publisher='Vervuert', publisher_place='Frankfurt am Main', collection_title='Lateinamerika-Studien', collection_number='43', number_of_pages='272'),
    seed([109], ACADEMIC, 'sammelbaende', ['zh'], 2021, 'periodical', title='洪堡研究', editor=[ETTE], container_title='外国语言与文化 / Foreign Languages and Cultures', volume='5', issue='3', page='1–74'),
    seed([115], ACADEMIC, 'didaktische-materialien', ['es'], 1993, 'book', title='Guillaume-Thomas Raynal, Alejandro de Humboldt y la imagen del hombre. Viajes y antropología a finales del siglo XVIII y principios del siglo XIX. Notas de apoyo a la docencia', author=[ETTE], publisher='Universidad Autónoma del Estado de México', publisher_place='Toluca', number_of_pages='75'),
    seed([120], ACADEMIC, 'hoerbuecher', ['de'], 2004, 'book', title='KOSMOS. Entwurf einer physischen Weltbeschreibung', author=[HUMBOLDT], editor=[ETTE, person('Lubrich', 'Oliver')], publisher='Eichborn', publisher_place='Frankfurt am Main', medium='Audio-CD', dimensions='2 Audio-CDs'),
    seed([127, 128], 'internet-praesentationen', None, ['de', 'en', 'es'], 1998, 'webpage', title='Alexander von Humboldt im Netz / Alexander von Humboldt in the net / Alexander von Humboldt en la red'),
    seed([176], 'aufsaetze', None, ['de'], 2000, 'article-journal', author=[ETTE], container_title='HiN – Alexander von Humboldt im Netz. Internationale Zeitschrift für Humboldt-Studien', volume='I', issue='1', page='5–28'),
    seed([238], 'aufsaetze', None, ['en'], 2006, 'article-journal', author=[person('Lubrich', 'Oliver'), ETTE], container_title='The Indexer – The International Journal of Indexing', volume='XXV', issue='1', page='2–6'),
    seed([308], 'aufsaetze', None, ['de'], 2011, 'chapter', author=[ETTE], container_title='Zentralamerika. Centroamérica', publisher='Editorial Universidad de Costa Rica', publisher_place='San José, Costa Rica', page='xvi und 2–16 (nur gerade Seiten)'),
    seed([408], 'aufsaetze', None, ['ru'], 2018, 'article-journal', title='Aleksandr fon Gumbolʹdt: kočujuščaja myslʹ i žiznennaja nauka', author=[ETTE], container_title='Epistemology & Philosophy of Science', volume='LV', issue='1', page='136–154'),
    seed([420], 'aufsaetze', None, ['zh'], 2019, 'article-journal', author=[ETTE], translator=[person('Ren', 'Haiyan')], container_title='外国语言与文化 / Foreign Languages and Cultures', volume='3', issue='3', page='2–11'),
    seed([444], 'aufsaetze', None, ['zh'], 2021, 'article-journal', author=[ETTE], translator=[person('Fan', 'Ni')], container_title='外国语言与文化 / Foreign Languages and Cultures', volume='5', issue='3', page='1–6'),
    seed([529], PREFACES, None, ['pt', 'de'], 2019, 'chapter', author=[ETTE], container_title='A Carta de Américo Vespúcio, em Lisboa (1504), e o mapa de Martin Waldseemüller, com o nome „America“, delineando as terras do futuro „Brasil“ (1507) / Amerigo Vespuccis Brief aus Lissabon (1504); wie Martin Waldseemüller den Namen „America“ erfand und damit das künftige „Brasilien“ meinte', publisher='Edições UFC', publisher_place='Fortaleza', page='10–13'),
    seed([531], PREFACES, None, ['zh'], 2019, 'article-journal', author=[ETTE], translator=[person('Ren', 'Haiyan')], container_title='外国语言与文化 / Foreign Languages and Cultures', volume='3', issue='3', page='1'),
    seed([548], 'debatten', None, ['de'], 2001, 'article-newspaper', author=[ETTE], container_title='PUTZ – Die Potsdamer Universitätszeitung', issue='3', page='25'),
    seed([553], 'uebersetzungen', None, ['de'], 1991, 'book', title='Reise in die Äquinoktial-Gegenden des Neuen Kontinents', author=[HUMBOLDT], publisher='Insel Verlag', publisher_place='Frankfurt am Main / Leipzig', number_of_volumes='2'),
    seed([557], 'uebersetzungen', None, ['de'], 2018, 'book', title='Das Buch der Begegnungen. Menschen – Kulturen – Geschichten aus den Amerikanischen Reisetagebüchern', author=[HUMBOLDT], editor=[ETTE], translator=[ETTE], publisher='Manesse Verlag', publisher_place='München', number_of_pages='395 + XVIII')
]

def source_html(paragraph):
    fragments = []
    for node in paragraph._p:
        if node.tag.endswith('}r'):
            value = ''.join(node.xpath('.//w:t/text()'))
            escaped = html.escape(value)
            italic = node.find('w:rPr/w:i', NAMESPACES)
            if italic is not None and italic.get('{' + NAMESPACES['w'] + '}val', '1') not in ['0', 'false', 'off']:
                escaped = '<em>' + escaped + '</em>'
            fragments.append(escaped)
        elif node.tag.endswith('}hyperlink'):
            value = ''.join(node.xpath('.//w:t/text()'))
            identifier = node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = paragraph.part.rels[identifier].target_ref
            assert target.startswith(('https://', 'http://'))
            fragments.append('<a href="' + html.escape(target, quote=True) + '">' + html.escape(value) + '</a>')
    result = ''.join(fragments)
    from lxml import etree
    recovered = ''.join(etree.fromstring(('<div>' + result + '</div>').encode()).itertext())
    assert recovered == paragraph.text
    return result

def extract_title(value):
    opening = value.index('„') + 1
    if '“ [' in value:
        return value[opening:value.index('“ [')]
    return value[opening:value.index('“. In:')]

def urls(value):
    return [match.rstrip('.') for match in re.findall(r'https?://\S+', value)]

def make_records(document):
    records = []
    category_map = {category['id']: category for category in CATEGORIES}
    for index, specification in enumerate(SEEDS, 1):
        identifier = f'ette-{index:06d}'
        numbers = specification['paragraphs']
        paragraphs = [document.paragraphs[number - 1] for number in numbers]
        primary = paragraphs[0].text
        assert primary.strip()
        bibliographic = {'id': identifier, 'type': specification['standard_type']}
        for key, value in specification['fields'].items():
            bibliographic[key.replace('_', '-')] = deepcopy(value)
        if 'title' not in bibliographic:
            bibliographic['title'] = extract_title(primary)
        bibliographic['issued'] = {'date-parts': [[specification['year']]]}
        languages = specification['languages']
        if len(languages) == 1:
            bibliographic['language'] = languages[0]
        for address in urls(primary):
            if address.startswith('https://doi.org/'):
                bibliographic['DOI'] = address.removeprefix('https://doi.org/')
            else:
                bibliographic['URL'] = address
        if numbers[0] == 127:
            bibliographic['URL'] = urls(paragraphs[1].text)[0]
        assessed = {'status': 'inferred', 'basis': 'Sprachzuordnung anhand des Titels; keine neue Volltextprüfung.'}
        if not languages:
            assessed = {'status': 'undetermined', 'basis': 'Die Sprache eines ganzen Sammelbandes ist aus seinem deutschen Titel nicht sicher bestimmbar.'}
        if numbers[0] == 127:
            assessed = {'status': 'explicit', 'basis': 'Quellenbeschreibung nennt eine dreisprachige Präsentation; der Titel benennt Deutsch, Englisch und Spanisch.'}
        if numbers[0] == 529:
            assessed['basis'] = 'Portugiesischer und deutscher Beitragstitel mit zweisprachigem Buchnachweis; Beitragssprachen bei der fachlichen Freigabe bestätigen.'
        record = {
            'id': identifier, 'category_id': specification['category'], 'subcategory_id': specification['subtype'],
            'languages': languages, 'language_assessment': assessed,
            'publication_year': specification['year'], 'year_basis': 'online-start' if numbers[0] == 127 else 'publication',
            'publication_medium': 'audio' if numbers[0] == 120 else 'website' if numbers[0] == 127 else 'text',
            'bibliographic': bibliographic, 'parallel_titles': [], 'container_contributors': {},
            'related_editions': [], 'relations': [], 'notes': [], 'review_status': 'draft',
            'source': {'file': SOURCE.name, 'sha256': EXPECTED_SHA256,
                'paragraph_numbers': numbers, 'paragraphs': [paragraph.text for paragraph in paragraphs],
                'citation_text': '\n'.join(paragraph.text for paragraph in paragraphs),
                'citation_html': '<br>'.join(source_html(paragraph) for paragraph in paragraphs)}
        }
        if numbers[0] in [109, 420, 444, 531]:
            parallel = primary.split('“ [', 1)[1].split('].', 1)[0]
            record['parallel_titles'].append({'language': 'en' if numbers[0] in [420, 444] else 'de', 'text': parallel, 'status': 'source-parallel-title'})
        if numbers[0] == 408:
            parallel = primary.split(' / ', 1)[1].split('“. In:', 1)[0]
            record['parallel_titles'].append({'language': 'en', 'text': parallel, 'status': 'source-parallel-title'})
            record['notes'].append('Russischer Titel in Quellenumschrift; ʹ bleibt U+02B9. Englischer Paralleltitel ist keine zweite Publikationssprache.')
        if numbers[0] == 36:
            for number, year, kind in [(37, 2005, 'Digitale Ausgabe'), (38, 2020, 'Zweite Auflage')]:
                value = document.paragraphs[number - 1].text
                related = {'publication_year': year, 'label': kind, 'source_paragraph': number, 'citation_text': value}
                if urls(value):
                    related['url'] = urls(value)[0]
                record['related_editions'].append(related)
            record['notes'].append('Die digitale URL gehört zur Ausgabe von 2005; Hauptjahr dieses Eintrags bleibt 2002.')
        if numbers[0] == 50:
            for number, year, kind in [(51, 1999, 'Jubiläumsausgabe und Neudruck laut Quelle'), (52, 2004, 'Dritte Auflage')]:
                record['related_editions'].append({'publication_year': year, 'label': kind, 'source_paragraph': number, 'citation_text': document.paragraphs[number - 1].text})
            record['relations'].append({'type': 'same-publication', 'target_id': 'ette-000019'})
        if numbers[0] == 553:
            record['relations'].append({'type': 'same-publication', 'target_id': 'ette-000004'})
            record['notes'].append('Separater redaktioneller Nachweis der Überarbeitung und teilweisen Neuübersetzung. Rollenname wird nicht ohne ausdrückliche Angabe im Quellenabsatz ergänzt.')
            bibliographic['note'] = 'Bd. 1: S. 11–701; Bd. 2: S. 703–1512 (Überarbeitung und teilweise Neuübersetzung).'
        if numbers[0] == 109:
            record['notes'].append('Redaktionell Sammelbände, bibliographisch Zeitschriftendossier. Herausgeberangabe aus der finalen Quelle übernommen; keine unabhängige Rollenprüfung.')
        if numbers[0] == 444:
            record['relations'].append({'type': 'part-of', 'target_id': 'ette-000006'})
            record['container_contributors'] = {'editor': [ETTE]}
        if numbers[0] == 308:
            record['container_contributors'] = {'author': [HUMBOLDT], 'editor': [person('Pérez Brignoli', 'Héctor')], 'translator': [person('Kruse', 'Silvia')]}
            record['notes'].append('Seiten nicht zusammenhängend: xvi sowie die geraden Seiten 2–16. Silvia Kruse ist im Nachweis Übersetzerin des Buches; keine Übertragung dieser Rolle auf den deutschen Beitrag.')
        if numbers[0] == 529:
            record['container_contributors'] = {'editor': [person('Bezerra', 'José Augusto'), person('Schwamborn', 'Ingrid')], 'translator': [person('Schwamborn', 'Ingrid')]}
            record['notes'].append('David Jacobson wird in der Quelle als Übersetzer des englischen Anhangs genannt, nicht als Übersetzer des hier erfassten Beitrags.')
        if numbers[0] == 120:
            record['notes'].append('CSL book mit medium Audio-CD; späterer RIS-Export benötigt eine geprüfte Hörbuch-Zuordnung.')
        assert specification['category'] in category_map
        records.append(record)
    return records

def build_schema():
    person_schema = {'type': 'object', 'required': ['family', 'given'], 'additionalProperties': False,
        'properties': {key: {'type': 'string'} for key in ['family', 'given', 'non-dropping-particle']}}
    bibliographic = {'type': 'object', 'required': ['id', 'type', 'title', 'issued'], 'additionalProperties': False, 'properties': {
        'id': {'type': 'string'}, 'type': {'enum': ['article-journal', 'article-newspaper', 'book', 'chapter', 'periodical', 'webpage']},
        'issued': {'type': 'object', 'required': ['date-parts'], 'additionalProperties': False, 'properties': {'date-parts': {'type': 'array', 'minItems': 1, 'maxItems': 1, 'items': {'type': 'array', 'minItems': 1, 'maxItems': 1, 'items': {'type': 'integer'}}}}}}}
    for key in ['title', 'container-title', 'publisher', 'publisher-place', 'collection-title', 'collection-number', 'number-of-pages', 'number-of-volumes', 'volume', 'issue', 'page', 'medium', 'dimensions', 'note', 'language', 'DOI', 'URL']:
        bibliographic['properties'][key] = {'type': 'string'}
    for role in ['author', 'editor', 'translator']:
        bibliographic['properties'][role] = {'type': 'array', 'items': person_schema}
    string_list = {'type': 'array', 'items': {'type': 'string'}}
    roles = {'type': 'object', 'additionalProperties': False, 'properties': {role: {'type': 'array', 'items': person_schema} for role in ['author', 'editor', 'translator']}}
    source = {'type': 'object', 'additionalProperties': False, 'required': ['file', 'sha256', 'paragraph_numbers', 'paragraphs', 'citation_text', 'citation_html'], 'properties': {
        'file': {'type': 'string'}, 'sha256': {'type': 'string', 'pattern': '^[a-f0-9]{64}$'},
        'paragraph_numbers': {'type': 'array', 'minItems': 1, 'uniqueItems': True, 'items': {'type': 'integer', 'minimum': 1}},
        'paragraphs': string_list, 'citation_text': {'type': 'string'}, 'citation_html': {'type': 'string'}}}
    properties = {'id': {'type': 'string', 'pattern': '^ette-[0-9]{6}$'}, 'category_id': {'enum': [category['id'] for category in CATEGORIES]},
        'subcategory_id': {'enum': [None] + [subcategory['id'] for category in CATEGORIES for subcategory in category['subcategories']]},
        'languages': {'type': 'array', 'uniqueItems': True, 'items': {'enum': list(LANGUAGES)}},
        'language_assessment': {'type': 'object', 'required': ['status', 'basis'], 'additionalProperties': False, 'properties': {'status': {'enum': ['explicit', 'inferred', 'undetermined']}, 'basis': {'type': 'string'}}},
        'publication_year': {'type': 'integer', 'minimum': 1400, 'maximum': 2200}, 'year_basis': {'enum': ['publication', 'online-start']},
        'publication_medium': {'enum': ['text', 'audio', 'website']}, 'bibliographic': bibliographic,
        'parallel_titles': {'type': 'array', 'items': {'type': 'object', 'required': ['language', 'text', 'status'], 'additionalProperties': False, 'properties': {'language': {'enum': list(LANGUAGES)}, 'text': {'type': 'string'}, 'status': {'const': 'source-parallel-title'}}}},
        'container_contributors': roles,
        'related_editions': {'type': 'array', 'items': {'type': 'object', 'additionalProperties': False, 'required': ['publication_year', 'label', 'source_paragraph', 'citation_text'], 'properties': {'publication_year': {'type': 'integer'}, 'label': {'type': 'string'}, 'source_paragraph': {'type': 'integer'}, 'citation_text': {'type': 'string'}, 'url': {'type': 'string'}}}},
        'relations': {'type': 'array', 'items': {'type': 'object', 'additionalProperties': False, 'required': ['type', 'target_id'], 'properties': {'type': {'enum': ['same-publication', 'part-of']}, 'target_id': {'type': 'string'}}}},
        'notes': string_list, 'review_status': {'enum': ['draft', 'reviewed']}, 'source': source}
    subcategories = [subcategory['id'] for category in CATEGORIES for subcategory in category['subcategories']]
    record = {'type': 'object', 'additionalProperties': False, 'required': list(properties), 'properties': properties,
        'allOf': [{'if': {'properties': {'category_id': {'const': ACADEMIC}}}, 'then': {'properties': {'subcategory_id': {'enum': subcategories}}}, 'else': {'properties': {'subcategory_id': {'type': 'null'}}}}]}
    category_schema = {'type': 'object', 'required': ['id', 'label', 'subcategories'], 'additionalProperties': False, 'properties': {
        'id': {'type': 'string'}, 'label': {'type': 'string'}, 'subcategories': {'type': 'array', 'items': {'type': 'object', 'required': ['id', 'label', 'source_label'], 'additionalProperties': False, 'properties': {key: {'type': 'string'} for key in ['id', 'label', 'source_label']}}}}}
    return {'$schema': 'https://json-schema.org/draft/2020-12/schema', 'title': 'AvH-Bibliographie-Ette Erfassungsprobe', 'type': 'object', 'additionalProperties': False,
        'required': ['schema_version', 'dataset', 'categories', 'records'], 'properties': {
            'schema_version': {'const': '0.1.0'}, 'dataset': {'type': 'object', 'additionalProperties': False, 'required': ['title', 'status', 'created', 'record_count', 'source_file', 'source_sha256'], 'properties': {
                'title': {'type': 'string'}, 'status': {'const': 'sample'}, 'created': {'type': 'string'}, 'record_count': {'type': 'integer'}, 'source_file': {'type': 'string'}, 'source_sha256': {'type': 'string', 'pattern': '^[a-f0-9]{64}$'}}},
            'categories': {'type': 'array', 'items': category_schema}, 'records': {'type': 'array', 'items': record}}}

def validate_instance(value, definition, location='$'):
    expected = definition.get('type')
    if expected:
        types = {'object': dict, 'array': list, 'string': str, 'integer': int, 'null': type(None)}
        assert isinstance(value, types[expected]), (location, expected)
        if expected == 'integer':
            assert not isinstance(value, bool), location
    if 'enum' in definition:
        assert value in definition['enum'], (location, value)
    if 'const' in definition:
        assert value == definition['const'], location
    if isinstance(value, dict):
        assert all(key in value for key in definition.get('required', [])), location
        properties = definition.get('properties', {})
        if definition.get('additionalProperties') is False:
            assert set(value) <= set(properties), location
        for key, item in value.items():
            if key in properties:
                validate_instance(item, properties[key], location + '.' + key)
    if isinstance(value, list):
        assert len(value) >= definition.get('minItems', 0), location
        assert len(value) <= definition.get('maxItems', len(value)), location
        if definition.get('uniqueItems'):
            assert len({json.dumps(item, sort_keys=True) for item in value}) == len(value), location
        if 'items' in definition:
            for index, item in enumerate(value):
                validate_instance(item, definition['items'], location + '[' + str(index) + ']')
    if isinstance(value, str) and 'pattern' in definition:
        assert re.search(definition['pattern'], value), location
    if isinstance(value, int):
        assert value >= definition.get('minimum', value), location
        assert value <= definition.get('maximum', value), location
    for condition in definition.get('allOf', []):
        matched = all(value.get(key) == rule['const'] for key, rule in condition['if']['properties'].items())
        validate_instance(value, condition['then'] if matched else condition['else'], location)

def validate(dataset, schema, document):
    validate_instance(dataset, schema)
    records = dataset['records']
    assert len(records) == dataset['dataset']['record_count'] == 20
    identifiers = {record['id'] for record in records}
    assert len(identifiers) == 20
    assert {record['category_id'] for record in records} == {category['id'] for category in CATEGORIES}
    assert {record['subcategory_id'] for record in records if record['subcategory_id']} == {'monographien', 'editionen', 'sammelbaende', 'didaktische-materialien', 'hoerbuecher'}
    for record in records:
        assert record['bibliographic']['id'] == record['id']
        assert record['bibliographic']['issued']['date-parts'][0][0] == record['publication_year']
        expected = [document.paragraphs[number - 1].text for number in record['source']['paragraph_numbers']]
        assert expected == record['source']['paragraphs']
        assert '\n'.join(expected) == record['source']['citation_text']
        for relation in record['relations']:
            assert relation['target_id'] in identifiers
    assert records[2]['publication_year'] == 2002
    assert 'URL' not in records[2]['bibliographic']
    assert [edition['publication_year'] for edition in records[2]['related_editions']] == [2005, 2020]
    assert records[12]['languages'] == ['ru']
    assert records[13]['languages'] == ['zh']
    assert records[14]['languages'] == ['zh']
    assert records[15]['languages'] == ['pt', 'de']
    assert records[4]['languages'] == []
    assert records[11]['bibliographic']['page'] == 'xvi und 2–16 (nur gerade Seiten)'
    assert 'translator' not in records[11]['bibliographic']

def write_csv(records):
    category_labels = {category['id']: category['label'] for category in CATEGORIES}
    subtype_labels = {subtype['id']: subtype['label'] for category in CATEGORIES for subtype in category['subcategories']}
    columns = ['id', 'kategorie', 'untergruppe', 'publikationstyp_csl', 'jahr', 'jahresgrundlage', 'sprachen', 'sprachzuordnung', 'titel', 'personen_mit_rollen', 'uebergeordneter_titel', 'seiten', 'doi', 'url', 'weitere_ausgaben_jahre', 'pruefstatus', 'quellabsaetze', 'literaturangabe']
    rows = []
    for record in records:
        metadata = record['bibliographic']
        roles = {key: metadata[key] for key in ['author', 'editor', 'translator'] if key in metadata}
        row = dict(zip(columns, [record['id'], category_labels[record['category_id']], subtype_labels.get(record['subcategory_id'], ''), metadata['type'], record['publication_year'], record['year_basis'], ';'.join(record['languages']), record['language_assessment']['status'], metadata['title'], json.dumps(roles, ensure_ascii=False), metadata.get('container-title', ''), metadata.get('page', ''), metadata.get('DOI', ''), metadata.get('URL', ''), ';'.join(str(edition['publication_year']) for edition in record['related_editions']), record['review_status'], ';'.join(map(str, record['source']['paragraph_numbers'])), record['source']['citation_text']]))
        rows.append(row)
    with (PROJECT / 'data/probe.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    with (PROJECT / 'data/probe.csv').open(encoding='utf-8-sig', newline='') as handle:
        loaded = list(csv.DictReader(handle))
    assert len(loaded) == 20
    assert [row['literaturangabe'] for row in loaded] == [record['source']['citation_text'] for record in records]

def write_html(records):
    navigation = ' · '.join('<a href="#' + category['id'] + '">' + html.escape(category['label']) + '</a>' for category in CATEGORIES)
    sections = []
    for category in CATEGORIES:
        entries = []
        for record in records:
            if record['category_id'] != category['id']:
                continue
            metadata = record['bibliographic']
            subtitle = next((subtype['label'] for subtype in category['subcategories'] if subtype['id'] == record['subcategory_id']), '')
            language = ', '.join(LANGUAGES[code] for code in record['languages']) or 'Sprache noch ungeklärt'
            labels = ' · '.join(filter(None, [record['id'], str(record['publication_year']), language, subtitle]))
            details = '<dl><dt>Bibliographischer Typ</dt><dd>' + html.escape(metadata['type']) + '</dd><dt>Sprachzuordnung</dt><dd>' + html.escape(record['language_assessment']['basis']) + '</dd><dt>Quellabsätze</dt><dd>' + ', '.join(map(str, record['source']['paragraph_numbers'])) + '</dd></dl>'
            details += ''.join('<p>' + html.escape(note) + '</p>' for note in record['notes'])
            entries.append('<article id="' + record['id'] + '"><h3>' + html.escape(metadata['title']) + '</h3><p class="labels">' + html.escape(labels) + '</p><p class="citation">' + record['source']['citation_html'] + '</p><details><summary>Erfassungsentscheidungen ansehen</summary>' + details + '</details></article>')
        sections.append('<section id="' + category['id'] + '"><h2>' + html.escape(category['label']) + '</h2>' + ''.join(entries) + '</section>')
    page = '<!doctype html><html lang="de"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>AvH-Bibliographie-Ette – Erfassungsprobe</title><style>body{font:17px/1.6 system-ui,sans-serif;color:#222;background:#fafafa;max-width:1000px;margin:auto;padding:24px}a{color:#075b8f}nav{padding:16px;background:#eee}h1{line-height:1.2}h2{border-bottom:2px solid #555;padding-top:24px}h3{font-size:1.08rem;line-height:1.4}article{background:white;border:1px solid #ddd;padding:20px;margin:16px 0}.labels{color:#555;font-size:.9rem}.citation{white-space:pre-wrap;overflow-wrap:anywhere}summary{cursor:pointer}dt{font-weight:600}dd{margin-left:0}a:focus-visible,summary:focus-visible{outline:3px solid #c66}footer{margin-top:40px;font-size:.9rem}</style></head><body><header><h1>AvH-Bibliographie-Ette</h1><p>Erfassungsprobe: 20 Einträge aus der finalen Word-Bibliographie. Entwurf zur fachlichen Prüfung, 13.10.2026.</p><p>Alle Angaben sind in dieser Datei enthalten. Die Ansicht funktioniert offline und ohne JavaScript. Externe Quellenlinks benötigen eine Internetverbindung. Textsuche: Strg+F.</p></header><nav aria-label="Redaktionelle Kategorien">' + navigation + '</nav><main>' + ''.join(sections) + '</main><footer>Quelle: bibliographie_ette_final.docx. Noch kein vollständiger Datensatz und keine fachliche Freigabe. JSON, CSV und Erfassungsregeln liegen im Projektordner.</footer></body></html>'
    assert '<script' not in page and 'stylesheet' not in page
    from lxml import html as html_parser
    tree = html_parser.fromstring(page)
    assert len(tree.xpath('//article')) == 20
    for record in records:
        element = tree.get_element_by_id(record['id']).xpath('./p[@class="citation"]')[0]
        rendered = ''.join(element.itertext())
        assert rendered == ''.join(record['source']['paragraphs'])
    (PROJECT / 'probe.html').write_text(page, encoding='utf-8')

def main():
    original = SOURCE.read_bytes()
    assert hashlib.sha256(original).hexdigest() == EXPECTED_SHA256, 'Die Quelle hat sich geändert; Absatzzuordnungen müssen erneut geprüft werden.'
    document = Document(SOURCE)
    records = make_records(document)
    dataset = {'schema_version': '0.1.0', 'dataset': {'title': 'AvH-Bibliographie-Ette', 'status': 'sample', 'created': DATE, 'record_count': len(records), 'source_file': SOURCE.name, 'source_sha256': EXPECTED_SHA256}, 'categories': CATEGORIES, 'records': records}
    schema = build_schema()
    validate(dataset, schema, document)
    if '--check' in sys.argv:
        existing = json.loads((PROJECT / 'data/probe.json').read_text(encoding='utf-8'))
        validate(existing, schema, document)
        assert existing == dataset, 'JSON weicht von der reproduzierbaren Erfassungsprobe ab.'
        assert json.loads((PROJECT / 'schema/dataset.schema.json').read_text(encoding='utf-8')) == schema
        print('Prüfung erfolgreich: Schema, 20 Einträge, Kategorien, Rollen, Jahreszuordnung, Beziehungen und Quellenidentität.')
        return
    for directory in ['data', 'schema']:
        (PROJECT / directory).mkdir(exist_ok=True)
    (PROJECT / 'data/probe.json').write_text(json.dumps(dataset, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (PROJECT / 'schema/dataset.schema.json').write_text(json.dumps(schema, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    write_csv(records)
    write_html(records)
    assert SOURCE.read_bytes() == original
    print(json.dumps({'records': len(records), 'categories': dict(Counter(record['category_id'] for record in records)), 'language_assessment': dict(Counter(record['language_assessment']['status'] for record in records)), 'source_unchanged': True}, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
